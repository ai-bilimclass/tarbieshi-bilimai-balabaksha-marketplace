#!/usr/bin/env python3
"""Deterministic DOCX builder for weekly cyclograms and daily plans.

The table structure is always cloned from the retained methodological DOCX.
Input content is supplied as JSON; the script refuses altered row labels,
unexpected columns, colored text, colored borders, and cell shading.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


FONT_NAME = "Times New Roman"
FONT_SIZE = Pt(12)
BLACK = "000000"
DAYS = ["Пн", "Вт", "Ср", "Чт", "Пт"]
DAY_ALIASES = {
    "пн": "Пн", "понедельник": "Пн",
    "вт": "Вт", "вторник": "Вт",
    "ср": "Ср", "среда": "Ср",
    "чт": "Чт", "четверг": "Чт",
    "пт": "Пт", "пятница": "Пт",
    "дүйсенбі": "Пн", "сейсенбі": "Вт", "сәрсенбі": "Ср",
    "бейсенбі": "Чт", "жұма": "Пт",
    "mon": "Пн", "monday": "Пн", "tue": "Вт", "tuesday": "Вт",
    "wed": "Ср", "wednesday": "Ср", "thu": "Чт", "thursday": "Чт",
    "fri": "Пт", "friday": "Пт",
}
LOCALIZATION = {
    "ru": {"weekly": "Недельная циклограмма", "daily": "План работы на день", "regime": "Примерный режим дня", "days": DAYS,
           "fields": ["Организация", "Группа", "Возраст", "Период", "Тема", "Ценность", "Дата"]},
    "kk": {"weekly": "Апталық циклограмма", "daily": "Күндік жұмыс жоспары", "regime": "Күннің үлгілік режимі",
           "days": ["Дүйсенбі", "Сейсенбі", "Сәрсенбі", "Бейсенбі", "Жұма"],
           "fields": ["Ұйым", "Топ", "Жас ерекшелігі", "Кезең", "Тақырып", "Құндылық", "Күні"]},
    "en": {"weekly": "Weekly cyclogram", "daily": "Daily work plan", "regime": "Approximate daily routine",
           "days": ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"],
           "fields": ["Organization", "Group", "Age", "Period", "Topic", "Value", "Date"]},
}
# The source template remains the canonical key for JSON validation. These labels
# are displayed in the document in the selected language, in the same row order.
DISPLAY_LABELS = {
    "kk": {
        "preschool": [
            "Балаларды қабылдау", "Ата-аналармен немесе баланың заңды өкілдерімен әңгімелесу, кеңес беру",
            "Балалардың іс-әрекеті (ойын, танымдық, коммуникативтік, шығармашылық, эксперименттік, еңбек, қимыл-қозғалыс, бейнелеу, дербес және басқа да түрлері)",
            "Таңертеңгі гимнастика", "Таңғы ас", "Ұйымдастырылған іс-әрекетке дайындық",
            "Кестеге сәйкес ұйымдастырылған іс-әрекет", "Екінші таңғы ас", "Серуенге дайындық", "Серуен",
            "Серуеннен оралу", "Түскі ас", "Күндізгі ұйқы", "Біртіндеп ояту, сауықтыру шаралары",
            "Балалардың дербес іс-әрекеті (қимылды, ұлттық, сюжетті-рөлдік, үстел үсті және басқа ойындар, бейнелеу іс-әрекеті, кітап қарау және басқалары)",
            "Бесін ас", "Балалармен жеке жұмыс", "Серуенге дайындық", "Серуен", "Серуеннен оралу",
            "Кешкі ас", "Балалардың дербес іс-әрекеті (қимылды, ұлттық, сюжетті-рөлдік, үстел үсті және басқа ойындар, бейнелеу іс-әрекеті, кітап қарау және басқалары)",
            "Балалардың үйге қайтуы",
        ],
        "preprimary": [
            "Балаларды қабылдау", "Ата-аналармен немесе баланың заңды өкілдерімен әңгімелесу, кеңес беру",
            "Дербес іс-әрекет (қимылды, ұлттық, сюжетті-рөлдік, үстел үсті және басқа ойындар, бейнелеу іс-әрекеті, кітап қарау және басқалары)",
            "Таңертеңгі гимнастика", "Ұйымдастырылған іс-әрекетке дайындық",
            "Кестеге сәйкес ұйымдастырылған іс-әрекет", "Серуенге дайындық", "Серуен",
            "Серуеннен оралу", "Балалармен жеке жұмыс",
            "Балалардың іс-әрекеті (ойын, танымдық, коммуникативтік, шығармашылық, эксперименттік, еңбек, қимыл-қозғалыс, бейнелеу, дербес және басқа да түрлері)",
            "Балалардың үйге қайтуы",
        ],
    },
    "en": {
        "preschool": [
            "Arrival of children", "Conversations and consultations with parents or legal guardians",
            "Children's activities (play, cognitive, communication, creative, experimental, work, movement, art, independent and other activities)",
            "Morning exercises", "Breakfast", "Preparation for organized activities", "Organized activities according to the schedule",
            "Second breakfast", "Preparation for outdoor time", "Outdoor time", "Return from outdoor time",
            "Lunch", "Afternoon nap", "Gradual waking and health activities",
            "Children's independent activities (active, national, role-play, tabletop and other games, art, looking at books and other activities)",
            "Afternoon snack", "Individual work with children", "Preparation for outdoor time", "Outdoor time",
            "Return from outdoor time", "Dinner",
            "Children's independent activities (active, national, role-play, tabletop and other games, art, looking at books and other activities)",
            "Departure of children",
        ],
        "preprimary": [
            "Arrival of children", "Conversations and consultations with parents or legal guardians",
            "Independent activities (active, national, role-play, tabletop and other games, art, looking at books and other activities)",
            "Morning exercises", "Preparation for organized activities", "Organized activities according to the schedule",
            "Preparation for outdoor time", "Outdoor time", "Return from outdoor time", "Individual work with children",
            "Children's activities (play, cognitive, communication, creative, experimental, work, movement, art, independent and other activities)",
            "Departure of children",
        ],
    },
}
WEEKLY_PARTS = {
    "цель": "goal", "мақсат": "goal", "goal": "goal",
    "задача": "task", "міндет": "task", "task": "task",
    "ход": "procedure", "барысы": "procedure", "procedure": "procedure",
    "ожидаемый результат": "result", "күтілетін нәтиже": "result", "expected result": "result",
}
WEEKLY_PART_PATTERN = re.compile(
    r"(?i)(?<!\w)(Цель|Мақсат|Goal|Задача|Міндет|Task|Ход|Барысы|Procedure|"
    r"Ожидаемый результат|Күтілетін нәтиже|Expected result)\s*:"
)
WALK_STAGE_PATTERN = re.compile(r"(?m)^\s*([1-4])\.\s+(.+)$")
WALK_DURATION_PATTERN = re.compile(
    r"(?i)(?:Длительность|Ұзақтығы|Duration)\s*:\s*(\d+)\s*(?:минут|minutes?)\b"
)
WALK_STAGE_TERMS = {
    "ru": (("наблюден",), ("подвижн", "спортивн"), ("труд",), ("самостоятельн",)),
    "kk": (("бақыла",), ("қимыл", "спорт"), ("еңбек",), ("дербес",)),
    "en": (("observation",), ("active", "sport"), ("work",), ("independent",)),
}
WALK_COGNITIVE_TERMS = {
    "ru": "познавательн",
    "kk": "таным",
    "en": "cognitive",
}


def _set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    for node in tr_pr.findall(qn("w:tblHeader")):
        tr_pr.remove(node)
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _set_cell_borders(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    old = tc_pr.find(qn("w:tcBorders"))
    if old is not None:
        tc_pr.remove(old)
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "8")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), BLACK)
        borders.append(element)
    tc_pr.append(borders)
    shading = tc_pr.find(qn("w:shd"))
    if shading is not None:
        tc_pr.remove(shading)


def _style_run(run, *, bold: bool | None = None) -> None:
    run.font.name = FONT_NAME
    run.font.size = FONT_SIZE
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), FONT_NAME)


def _style_paragraph(paragraph, *, bold: bool = False, centered: bool = False) -> None:
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    if centered:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if not paragraph.runs:
        paragraph.add_run("")
    for run in paragraph.runs:
        _style_run(run, bold=bold)


def _replace_cell_text(cell, text: str, *, bold: bool = False, centered: bool = False) -> None:
    cell.text = str(text or "")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        _style_paragraph(paragraph, bold=bold, centered=centered)
    _set_cell_borders(cell)


def _remove_blank_regime_rows(table) -> None:
    for row in list(table.rows)[1:]:
        if not row.cells[0].text.strip():
            table._tbl.remove(row._tr)


def _remove_columns_except(table, keep: list[int]) -> None:
    keep_set = set(keep)
    grid = table._tbl.tblGrid
    for index in reversed(range(len(grid.gridCol_lst))):
        if index not in keep_set:
            grid.remove(grid.gridCol_lst[index])
    for row in table.rows:
        cells = row._tr.tc_lst
        for index in reversed(range(len(cells))):
            if index not in keep_set:
                row._tr.remove(cells[index])


def _normalize_day(value: str) -> str:
    key = (value or "").strip().lower()
    result = DAY_ALIASES.get(key)
    if not result:
        raise ValueError("Для плана дня поле day должно обозначать Пн, Вт, Ср, Чт или Пт.")
    return result


def _validate_weekly_content(value: object, label: str, day: str) -> None:
    text = str(value or "").strip()
    matches = list(WEEKLY_PART_PATTERN.finditer(text))
    parts = [WEEKLY_PARTS[match.group(1).casefold()] for match in matches]
    if parts != ["goal", "task", "procedure", "result"]:
        raise ValueError(
            f"Строка «{label}», {day}: нужны четыре части в порядке "
            "Цель/Мақсат, Задача/Міндет, Ход/Барысы, "
            "Ожидаемый результат/Күтілетін нәтиже."
        )
    content = [
        text[match.end():matches[index + 1].start() if index + 1 < len(matches) else len(text)].strip()
        for index, match in enumerate(matches)
    ]
    organized = "организованн" in label.casefold()
    minimums = [20, 30, 190 if organized else 110, 25]
    if any(len(part) < minimum for part, minimum in zip(content, minimums)):
        raise ValueError(
            f"Строка «{label}», {day}: формулировки слишком краткие. "
            "Раскрой конкретную задачу, последовательность действий педагога "
            "и детей с материалом и наблюдаемый результат."
        )


def _validate_walk_content(value: object, label: str, location: str, language: str) -> None:
    if label.strip().casefold() != "прогулка":
        return
    text = str(value or "")
    stages = list(WALK_STAGE_PATTERN.finditer(text))
    terms = WALK_STAGE_TERMS[language]
    if [match.group(1) for match in stages] != ["1", "2", "3", "4"] or any(
        not any(term in match.group(2).casefold() for term in alternatives)
        for alternatives, match in zip(terms, stages)
    ):
        raise ValueError(
            f"Строка «{label}», {location}: у прогулки нужны четыре нумерованные части "
            "в порядке: наблюдение в природе, подвижная игра, труд в природе, "
            "самостоятельная деятельность детей."
        )
    duration = WALK_DURATION_PATTERN.search(text)
    if duration is None or int(duration.group(1)) < 90:
        raise ValueError(
            f"Строка «{label}», {location}: явно укажи длительность прогулки "
            "не менее 90 минут на языке документа."
        )
    if WALK_COGNITIVE_TERMS[language] not in text.casefold():
        raise ValueError(
            f"Строка «{label}», {location}: раскрой познавательную деятельность "
            "в ходе наблюдения на прогулке."
        )


def _validate_rows(data: dict, expected_labels: list[str], mode: str, language: str) -> list[dict]:
    rows = data.get("rows")
    if not isinstance(rows, list):
        raise ValueError("Поле rows должно быть массивом в точном порядке строк шаблона.")
    labels = [str(item.get("label", "")).strip() for item in rows if isinstance(item, dict)]
    if labels != expected_labels:
        raise ValueError(
            "Названия или порядок режимных строк отличаются от шаблона. "
            "Запрещено добавлять, удалять, объединять, переименовывать или переставлять строки."
        )
    for item in rows:
        if mode == "weekly":
            days = item.get("days")
            if not isinstance(days, dict) or set(days) != set(DAYS):
                raise ValueError(f"В каждой недельной строке должны быть ровно ключи: {', '.join(DAYS)}.")
            for day in DAYS:
                _validate_weekly_content(days[day], item["label"], day)
                _validate_walk_content(days[day], item["label"], day, language)
        elif "content" not in item:
            raise ValueError("В каждой строке плана дня требуется поле content.")
        else:
            _validate_walk_content(item["content"], item["label"], "план дня", language)
    return rows


def _copy_page_geometry(source, target) -> None:
    src = source.sections[0]
    dst = target.sections[0]
    dst.orientation = WD_ORIENT.LANDSCAPE
    dst.page_width = max(src.page_width, src.page_height)
    dst.page_height = min(src.page_width, src.page_height)
    dst.left_margin = src.left_margin
    dst.right_margin = src.right_margin
    dst.top_margin = src.top_margin
    dst.bottom_margin = src.bottom_margin
    dst.header_distance = src.header_distance
    dst.footer_distance = src.footer_distance


def _add_metadata(document, data: dict, mode: str, language: str) -> None:
    local = LOCALIZATION[language]
    title = local[mode]
    p = document.add_paragraph()
    run = p.add_run(title)
    _style_run(run, bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    fields = list(zip(local["fields"][:6], [
        data.get("organization", "____________________________"), data.get("group", ""),
        data.get("age", ""), data.get("period", ""), data.get("topic", ""), data.get("value", ""),
    ]))
    if mode == "daily":
        fields.insert(4, (local["fields"][6], data.get("date", "")))
    for label, value in fields:
        paragraph = document.add_paragraph()
        label_run = paragraph.add_run(f"{label}: ")
        value_run = paragraph.add_run(str(value or ""))
        _style_run(label_run, bold=True)
        _style_run(value_run, bold=False)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)


def _append_table(document, template_table):
    document._body._body.append(deepcopy(template_table._tbl))
    return document.tables[-1]


def _audit(document, table, expected_labels: list[str], mode: str) -> None:
    if any(section.orientation != WD_ORIENT.LANDSCAPE or section.page_width <= section.page_height
           for section in document.sections):
        raise AssertionError("Все страницы итогового DOCX должны иметь альбомную ориентацию.")
    expected_columns = 6 if mode == "weekly" else 2
    if len(table.columns) != expected_columns:
        raise AssertionError(f"Ожидалось {expected_columns} столбца(ов), получено {len(table.columns)}.")
    actual_labels = [row.cells[0].text.strip() for row in table.rows[1:]]
    if actual_labels != expected_labels:
        raise AssertionError("Первый столбец итогового документа не совпадает с шаблоном.")
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            if tc_pr.find(qn("w:shd")) is not None:
                raise AssertionError("Обнаружена запрещённая цветная заливка ячейки.")
            borders = tc_pr.find(qn("w:tcBorders"))
            if borders is None or any(edge.get(qn("w:color")) != BLACK for edge in borders):
                raise AssertionError("Все границы таблицы должны быть чёрными.")
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    if run.font.name != FONT_NAME or run.font.size != FONT_SIZE:
                        raise AssertionError("В таблице допускается только Times New Roman 12 pt.")
                    if run.font.color.rgb != RGBColor(0, 0, 0):
                        raise AssertionError("В таблице допускается только чёрный текст.")
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            if run.font.name != FONT_NAME or run.font.size != FONT_SIZE:
                raise AssertionError("В документе допускается только Times New Roman 12 pt.")
            if run.font.color.rgb != RGBColor(0, 0, 0):
                raise AssertionError("В документе допускается только чёрный текст.")


def build(args: argparse.Namespace) -> None:
    data = json.loads(Path(args.data).read_text(encoding="utf-8"))
    language = str(data.get("language", "")).strip().lower()
    if language not in LOCALIZATION:
        raise ValueError("Поле language обязательно: kk, ru или en.")
    source = Document(args.template)
    table_index = 1 if args.group_type == "preschool" else 2
    source_table = source.tables[table_index]

    output = Document()
    _copy_page_geometry(source, output)
    _add_metadata(output, data, args.mode, language)
    table = _append_table(output, source_table)
    _remove_blank_regime_rows(table)
    expected_labels = [row.cells[0].text.strip() for row in table.rows[1:]]
    expected_count = 23 if args.group_type == "preschool" else 12
    if len(expected_labels) != expected_count:
        raise AssertionError(
            f"Шаблон содержит {len(expected_labels)} непустых режимных строк вместо {expected_count}."
        )
    rows = _validate_rows(data, expected_labels, args.mode, language)
    display_labels = expected_labels if language == "ru" else DISPLAY_LABELS[language][args.group_type]
    if len(display_labels) != expected_count:
        raise AssertionError("Число переведённых названий строк не совпадает с формой шаблона.")

    if args.mode == "weekly":
        headers = [LOCALIZATION[language]["regime"], *LOCALIZATION[language]["days"]]
        for col, text in enumerate(headers):
            _replace_cell_text(table.cell(0, col), text, bold=True, centered=True)
        for row_index, item in enumerate(rows, start=1):
            _replace_cell_text(table.cell(row_index, 0), display_labels[row_index - 1], bold=True)
            for col_index, day in enumerate(DAYS, start=1):
                _replace_cell_text(table.cell(row_index, col_index), item["days"][day])
    else:
        day = _normalize_day(data.get("day", ""))
        day_index = DAYS.index(day) + 1
        _remove_columns_except(table, [0, day_index])
        date = str(data.get("date", "")).strip()
        second_header = f"{LOCALIZATION[language]['days'][DAYS.index(day)]} {date}".strip()
        _replace_cell_text(table.cell(0, 0), LOCALIZATION[language]["regime"], bold=True, centered=True)
        _replace_cell_text(table.cell(0, 1), second_header, bold=True, centered=True)
        for row_index, item in enumerate(rows, start=1):
            _replace_cell_text(table.cell(row_index, 0), display_labels[row_index - 1], bold=True)
            _replace_cell_text(table.cell(row_index, 1), item["content"])

    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_repeat_header(table.rows[0])
    for row in table.rows:
        for cell in row.cells:
            _set_cell_borders(cell)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    _style_run(run)
    for run in table.rows[0].cells[0].paragraphs[0].runs:
        run.bold = True
    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True

    _audit(output, table, display_labels, args.mode)
    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output.save(output_path)
    reopened = Document(output_path)
    _audit(reopened, reopened.tables[-1], display_labels, args.mode)


def parse_args() -> argparse.Namespace:
    plugin_root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser()
    parser.add_argument("--mode", choices=("weekly", "daily"), required=True)
    parser.add_argument("--group-type", choices=("preschool", "preprimary"), required=True)
    parser.add_argument("--data", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument(
        "--template",
        default=str(plugin_root / "references" / "Методическое сопровождение для Плагина ДО.docx"),
    )
    return parser.parse_args()


if __name__ == "__main__":
    try:
        build(parse_args())
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        sys.exit(1)
