#!/usr/bin/env python3
"""Build an editable preschool weekly kit in DOCX from structured JSON.

The builder is topic-neutral: text remains editable Word content, while optional
illustrations are embedded from paths supplied in the JSON. Run with
``--write-example week.json`` to create a complete example input.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


COLORS = {
    "green": "26735B",
    "light_green": "EAF5EF",
    "gold": "E5A92E",
    "red": "C74A3D",
    "text": "24302C",
    "white": "FFFFFF",
}

AGE_GROUPS = {"1–3", "3–4", "4–5", "5–6", "6–7", "1-3", "3-4", "4-5", "5-6", "6-7"}
DAYS = ["Понедельник", "Вторник", "Среда", "Четверг", "Пятница"]

EXAMPLE_INPUT: dict[str, Any] = {
    "title": "Редактируемый недельный комплект",
    "age": "4–5",
    "theme": "Транспорт",
    "priorities": ["речевое развитие", "логическое мышление", "мелкая моторика"],
    "accent": "счёт до пяти",
    "summary": "Неделя знакомит детей с назначением транспорта через речь, движение, логику и творчество.",
    "methodical_rationale": "Нагрузка чередуется с движением и продуктивной деятельностью. Новые слова закрепляются в игре и итоговом сценарии.",
    "week": [
        {"day": day, "episodes": [
            {"time": "Утро, 10 мин", "activity": "Беседа", "content": f"Рассматриваем транспорт: {day.lower()}."},
            {"time": "До прогулки, 15 мин", "activity": "Организованная деятельность", "content": "Сравниваем транспорт по назначению и считаем до пяти."},
            {"time": "Прогулка, 90 мин", "activity": "Прогулка", "content": "1. Наблюдение в природе: обсуждаем погодные условия и проводим познавательное сравнение. 2. Подвижная игра «Светофор»: движение по цветовым сигналам. 3. Труд в природе: посильный уход за участком. 4. Самостоятельная деятельность детей: игры малой подвижности и сюжетно-ролевые игры по выбору."},
            {"time": "После сна, 10 мин", "activity": "Закрепление", "content": "Выбираем подходящий транспорт для ситуации."},
        ]} for day in DAYS
    ],
    "cards": [
        {"label": "Автобус"},
        {"label": "Поезд"},
        {"label": "Самолёт"},
        {"label": "Корабль"},
        {"label": "Велосипед"},
    ],
    "pair_tasks": [
        {"title": "Кому где двигаться?", "instruction": "Соедини транспорт и место движения.", "items": ["корабль — вода", "поезд — рельсы", "самолёт — небо"]}
    ],
    "logic_tasks": [
        {"title": "Найди лишнее", "instruction": "Покажи предмет, который не перевозит людей.", "images": [], "shown": "автобус, поезд, чашка", "answer": "чашка", "help": "Назвать назначение каждого предмета.", "age_hint": "Оставить три изображения."}
    ],
    "graphomotor": {"instruction": "Проведи транспорт по дорожкам.", "patterns": ["────────────", "∿∿∿∿∿∿∿∿", "○○○○○○○○", "/\\/\\/\\/\\", "﹏﹏﹏﹏﹏﹏", "╱ ╱ ╱ ╱ ╱ ╱"]},
    "coloring": {"title": "Раскраска по теме недели", "instruction": "Раскрась крупные части изображения."},
    "crafts": [{
        "title": "Аппликация «Автобус»",
        "type": "аппликация",
        "materials": ["крупные бумажные детали", "клей-карандаш", "лист-основа"],
        "steps": ["Рассмотреть крупные детали.", "Приклеить корпус.", "Добавить окна и колёса.", "Рассмотреть готовую работу."],
        "result_description": "Крупный автобус из простых геометрических деталей."
    }],
    "final_scenario": ["Приветствие и загадка о транспорте.", "Выбор карточек по назначению.", "Подвижная игра «Светофор».", "Показ творческих работ и речевое закрепление."],
    "extras": ["Наблюдение на прогулке: какой транспорт проезжает рядом.", "Словарь недели: пассажир, водитель, остановка, маршрут."],
}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def clear_cell_shading(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for shd in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(shd)


def set_cell_border(cell, color: str = COLORS["green"], size: str = "10") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_margins(cell, value: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "left", "bottom", "right"):
        elem = tc_mar.find(qn(f"w:{side}"))
        if elem is None:
            elem = OxmlElement(f"w:{side}")
            tc_mar.append(elem)
        elem.set(qn("w:w"), str(value))
        elem.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.add_run("Страница ")
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)


def configure_section(section, landscape: bool = False, planning: bool = False) -> None:
    section.page_width = Mm(297 if landscape else 210)
    section.page_height = Mm(210 if landscape else 297)
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.top_margin = Mm(17)
    section.bottom_margin = Mm(17)
    section.left_margin = Mm(17)
    section.right_margin = Mm(17)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    header = section.header.paragraphs[0]
    header.clear()
    header.text = "Tarbieshi BilimAI Balabaksha — редактируемый комплект"
    header.style = "Caption"
    footer = section.footer.paragraphs[0]
    footer.clear()
    add_page_field(footer)
    if planning:
        style_planning_paragraph(header)
        style_planning_paragraph(footer)


def set_run_typeface(run, name: str, size: int, color: str) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._r.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        r_fonts.set(qn(attr), name)


def style_planning_paragraph(paragraph) -> None:
    for run in paragraph.runs:
        set_run_typeface(run, "Times New Roman", 12, "000000")


def style_planning_table(table, header: bool = True) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_border(cell, color="000000", size="4")
            clear_cell_shading(cell)
            for paragraph in cell.paragraphs:
                style_planning_paragraph(paragraph)
                if header and row_index == 0:
                    for run in paragraph.runs:
                        run.font.bold = True
        if header and row_index == 0:
            set_repeat_table_header(row)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(COLORS["text"])
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for name, size in (("Title", 26), ("Heading 1", 19), ("Heading 2", 15), ("Heading 3", 12)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(COLORS["green"])
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)
    doc.styles["Caption"].font.name = "Arial"
    doc.styles["Caption"].font.size = Pt(9)
    doc.styles["Caption"].font.color.rgb = RGBColor.from_string(COLORS["green"])


def style_table(table, header: bool = True) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_border(cell)
            if header and row_index == 0:
                set_cell_shading(cell, COLORS["green"])
                for run in cell.paragraphs[0].runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.bold = True
        if header and row_index == 0:
            set_repeat_table_header(row)


def add_heading_page(doc: Document, heading: str) -> None:
    doc.add_page_break()
    doc.add_heading(heading, level=1)


def add_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        doc.add_paragraph(str(item), style="List Bullet")


def resolve_image(raw: str | None, base_dir: Path) -> Path | None:
    if not raw:
        return None
    path = Path(raw).expanduser()
    if not path.is_absolute():
        path = base_dir / path
    path = path.resolve()
    if not path.is_file():
        raise FileNotFoundError(f"Image not found: {path}")
    return path


def add_image_or_placeholder(cell, raw: str | None, base_dir: Path, width_cm: float, placeholder: str = "Место для изображения") -> None:
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image = resolve_image(raw, base_dir)
    if image:
        paragraph.add_run().add_picture(str(image), width=Cm(width_cm))
    else:
        run = paragraph.add_run(f"[{placeholder}]")
        run.font.color.rgb = RGBColor.from_string(COLORS["green"])
        run.font.size = Pt(12)


def validate_data(data: dict[str, Any]) -> None:
    missing = [key for key in ("age", "theme", "week") if not data.get(key)]
    if missing:
        raise ValueError("Missing required fields: " + ", ".join(missing))
    if str(data["age"]).replace(" лет", "") not in AGE_GROUPS:
        raise ValueError("age must be one of 1–3, 3–4, 4–5, 5–6, 6–7")
    if len(data["week"]) != 5:
        raise ValueError("week must contain exactly five day objects")
    for day in data["week"]:
        if not 3 <= len(day.get("episodes", [])) <= 4:
            raise ValueError("each day must contain 3–4 episodes")


def add_cover_and_summary(doc: Document, data: dict[str, Any]) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    run = p.add_run(data.get("title", "Редактируемый недельный комплект"))
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor.from_string(COLORS["green"])
    for label, value in (("Возраст", data["age"]), ("Лексическая тема", data["theme"]), ("Дополнительный акцент", data.get("accent", "не задан"))):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(f"{label}: ").bold = True
        p.add_run(str(value))
    doc.add_page_break()
    doc.add_heading("Короткое резюме", level=1)
    doc.add_paragraph(data.get("summary", "Комплект объединяет планирование, печатные задания и творческие материалы."))
    doc.add_heading("Приоритетные навыки", level=2)
    add_bullets(doc, data.get("priorities") or ["сбалансированный возрастной набор"])


def add_weekly_plan(doc: Document, data: dict[str, Any]) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(section, landscape=True, planning=True)
    heading = doc.add_heading("Циклограмма", level=1)
    style_planning_paragraph(heading)
    week = data["week"]
    max_episodes = max(len(day["episodes"]) for day in week)
    table = doc.add_table(rows=max_episodes + 1, cols=6)
    table.columns[0].width = Cm(4.2)
    for col in range(1, 6):
        table.columns[col].width = Cm(4.45)
    headers = ["Период / вид / содержание"] + [str(day.get("day") or DAYS[i]) for i, day in enumerate(week)]
    for col, text in enumerate(headers):
        table.cell(0, col).text = text
    for row in range(max_episodes):
        table.cell(row + 1, 0).text = f"Эпизод {row + 1}"
        for col, day in enumerate(week, start=1):
            episodes = day["episodes"]
            if row >= len(episodes):
                continue
            episode = episodes[row]
            cell = table.cell(row + 1, col)
            cell.text = ""
            p = cell.paragraphs[0]
            p.add_run(str(episode.get("time", "")) + "\n").bold = True
            p.add_run(str(episode.get("activity", "")) + "\n").italic = True
            p.add_run(str(episode.get("content", "")))
    style_planning_table(table)
    rationale = doc.add_paragraph(data.get("methodical_rationale", "Направления равномерно распределены и чередуются с двигательной активностью."))
    style_planning_paragraph(rationale)


def add_cards(doc: Document, cards: list[dict[str, Any]], base_dir: Path, *, new_page: bool = True) -> None:
    if new_page:
        add_heading_page(doc, "Карточки по теме")
    else:
        doc.add_heading("Карточки по теме", level=1)
    if not cards:
        doc.add_paragraph("Карточки не включены в исходные данные.")
        return
    for page_start in range(0, len(cards), 4):
        if page_start:
            doc.add_page_break()
            doc.add_heading("Карточки по теме — продолжение", level=2)
        table = doc.add_table(rows=2, cols=2)
        table.autofit = False
        for offset in range(4):
            cell = table.cell(offset // 2, offset % 2)
            set_cell_border(cell, size="14")
            set_cell_margins(cell, 180)
            index = page_start + offset
            if index >= len(cards):
                cell.text = ""
                continue
            card = cards[index]
            add_image_or_placeholder(cell, card.get("image"), base_dir, 6.2)
            p = cell.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(card.get("label", "Карточка")))
            run.bold = True
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor.from_string(COLORS["text"])


def add_pair_tasks(doc: Document, tasks: list[dict[str, Any]]) -> None:
    add_heading_page(doc, "Парные и пространственные задания")
    if not tasks:
        doc.add_paragraph("Раздел не включён в исходные данные.")
    for task in tasks:
        doc.add_heading(str(task.get("title", "Задание")), level=2)
        doc.add_paragraph(str(task.get("instruction", "Покажи или соедини подходящие изображения.")))
        add_bullets(doc, task.get("items", []))


def add_logic_tasks(doc: Document, tasks: list[dict[str, Any]], base_dir: Path) -> None:
    add_heading_page(doc, "Логические карточки для детей")
    if not tasks:
        doc.add_paragraph("Раздел не включён в исходные данные.")
    for task in tasks:
        doc.add_heading(str(task.get("title", "Логическое задание")), level=2)
        doc.add_paragraph(str(task.get("instruction", "Покажи правильный вариант.")))
        images = task.get("images", [])
        if images:
            table = doc.add_table(rows=1, cols=len(images))
            for index, raw in enumerate(images):
                set_cell_border(table.cell(0, index))
                add_image_or_placeholder(table.cell(0, index), raw, base_dir, 4.0)
    add_heading_page(doc, "Письменные ключи для воспитателя")
    for index, task in enumerate(tasks, start=1):
        doc.add_heading(f"Задание {index}: {task.get('title', 'Логическое задание')}", level=2)
        for label, key in (("Что изображено", "shown"), ("Задание", "instruction"), ("Правильный ответ", "answer"), ("Допустимая помощь", "help"), ("Возрастная подсказка", "age_hint")):
            p = doc.add_paragraph()
            p.add_run(label + ": ").bold = True
            p.add_run(str(task.get(key, "—")))


def add_graphomotor(doc: Document, graph: dict[str, Any]) -> None:
    add_heading_page(doc, "Графомоторный рабочий лист")
    doc.add_paragraph(str(graph.get("instruction", "Проведи по линиям слева направо.")))
    patterns = graph.get("patterns") or ["────────────", "∿∿∿∿∿∿∿∿", "/\\/\\/\\/\\"]
    for pattern in patterns:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(pattern))
        run.font.name = "Arial"
        run.font.size = Pt(26)
        p.paragraph_format.space_after = Pt(16)
    add_heading_page(doc, "Вырезные графомоторные полоски")
    table = doc.add_table(rows=len(patterns), cols=1)
    for index, pattern in enumerate(patterns):
        cell = table.cell(index, 0)
        set_cell_border(cell, color=COLORS["gold"], size="14")
        set_cell_margins(cell, 180)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(pattern))
        run.font.size = Pt(24)


def add_coloring(doc: Document, coloring: dict[str, Any], base_dir: Path) -> None:
    add_heading_page(doc, str(coloring.get("title", "Раскраска")))
    doc.add_paragraph(str(coloring.get("instruction", "Раскрась крупные области.")))
    table = doc.add_table(rows=1, cols=1)
    set_cell_border(table.cell(0, 0), color=COLORS["text"], size="14")
    add_image_or_placeholder(table.cell(0, 0), coloring.get("image"), base_dir, 14.5, "Место для контурной раскраски")


def add_crafts(doc: Document, crafts: list[dict[str, Any]], base_dir: Path) -> None:
    add_heading_page(doc, "Письменные алгоритмы творчества")
    if not crafts:
        doc.add_paragraph("Раздел не включён в исходные данные.")
    for craft in crafts:
        doc.add_heading(str(craft.get("title", "Творческая работа")), level=2)
        doc.add_paragraph("Материалы:", style="Heading 3")
        add_bullets(doc, craft.get("materials", []))
        for index, step in enumerate(craft.get("steps", [])[:4], start=1):
            p = doc.add_paragraph()
            p.add_run(f"Шаг {index}. ").bold = True
            p.add_run(str(step))
        warning = doc.add_paragraph()
        warning.add_run("Безопасность: ").bold = True
        warning.add_run("Работа проводится при постоянном наблюдении взрослого.")
    add_heading_page(doc, "Наглядные алгоритмы")
    for craft in crafts:
        doc.add_heading(str(craft.get("title", "Творческая работа")), level=2)
        table = doc.add_table(rows=2, cols=2)
        images = list(craft.get("step_images", []))[:4]
        images += [""] * (4 - len(images))
        for index in range(4):
            cell = table.cell(index // 2, index % 2)
            set_cell_border(cell)
            add_image_or_placeholder(cell, images[index], base_dir, 6.2, f"Иллюстрация шага {index + 1}")
            p = cell.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(f"Шаг {index + 1}").bold = True
        doc.add_heading("Итоговый результат", level=3)
        result = doc.add_table(rows=1, cols=1).cell(0, 0)
        set_cell_border(result, color=COLORS["gold"])
        add_image_or_placeholder(result, craft.get("result_image"), base_dir, 9.0, "Изображение результата")
        result.add_paragraph(str(craft.get("result_description", "Готовая работа.")))


def add_scenario_and_extras(doc: Document, data: dict[str, Any]) -> None:
    add_heading_page(doc, "Итоговый сценарий недели")
    for index, item in enumerate(data.get("final_scenario", []), start=1):
        p = doc.add_paragraph(style="List Number")
        p.add_run(str(item))
    add_heading_page(doc, "Дополнительные материалы")
    add_bullets(doc, data.get("extras", ["Дополнительные материалы не заданы."]))


def build_docx(data: dict[str, Any], output: Path, base_dir: Path) -> Path:
    validate_data(data)
    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0], landscape=False)
    add_cover_and_summary(doc, data)
    add_weekly_plan(doc, data)
    third = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(third, landscape=False)
    add_cards(doc, list(data.get("cards", [])), base_dir, new_page=False)
    add_pair_tasks(doc, list(data.get("pair_tasks", [])))
    add_logic_tasks(doc, list(data.get("logic_tasks", [])), base_dir)
    add_graphomotor(doc, dict(data.get("graphomotor", {})))
    add_coloring(doc, dict(data.get("coloring", {})), base_dir)
    add_crafts(doc, list(data.get("crafts", [])), base_dir)
    add_scenario_and_extras(doc, data)
    output = output.expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, help="UTF-8 JSON file with weekly-kit data")
    parser.add_argument("--output", type=Path, help="Target .docx path")
    parser.add_argument("--write-example", type=Path, metavar="PATH", help="Write example JSON and exit")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if args.write_example:
        target = args.write_example.expanduser().resolve()
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(json.dumps(EXAMPLE_INPUT, ensure_ascii=False, indent=2), encoding="utf-8")
        print(target)
        return
    if not args.input or not args.output:
        raise SystemExit("Provide --input and --output, or use --write-example PATH")
    input_path = args.input.expanduser().resolve()
    data = json.loads(input_path.read_text(encoding="utf-8"))
    output = build_docx(data, args.output, input_path.parent)
    print(output)


if __name__ == "__main__":
    main()
