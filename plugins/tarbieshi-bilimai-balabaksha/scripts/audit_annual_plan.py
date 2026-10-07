#!/usr/bin/env python3
"""Check that each annual-plan table row has usable pedagogical detail."""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


SECTION_NAMES = {
    "цель": "goal", "мақсат": "goal", "goal": "goal",
    "задачи": "tasks", "міндеттер": "tasks", "tasks": "tasks",
    "основные виды деятельности и ход работы": "activities",
    "негізгі әрекеттер және жұмыс барысы": "activities",
    "main activities and procedure": "activities",
    "ожидаемый результат": "result",
    "күтілетін нәтиже": "result",
    "expected result": "result",
}
SECTION_RE = re.compile(
    r"(?im)^\s*(Цель|Мақсат|Goal|Задачи|Міндеттер|Tasks|"
    r"Основные виды деятельности и ход работы|"
    r"Негізгі әрекеттер және жұмыс барысы|Main activities and procedure|"
    r"Ожидаемый результат|Күтілетін нәтиже|Expected result)\s*:"
)
ITEM_RE = re.compile(r"(?m)^\s*[1-3][.)]\s+")
ACTIVITY_NAMES = {
    "название": "name", "атауы": "name", "name": "name",
    "материалы": "materials", "материалдар": "materials", "materials": "materials",
    "ход": "procedure", "барысы": "procedure", "procedure": "procedure",
}
ACTIVITY_RE = re.compile(
    r"(?i)(?<!\w)(Название|Атауы|Name|Материалы|Материалдар|Materials|"
    r"Ход|Барысы|Procedure)\s*:"
)
ADULT_RE = re.compile(
    r"(?i)педагог|воспитател|руководител|инструктор|взросл|"
    r"тәрбиеші|мұғалім|teacher|educator|instructor|adult"
)
CHILD_RE = re.compile(r"(?i)дети|реб[её]нок|ребята|балалар|бала|children|child")
BOILERPLATE = (
    "в начале месяца педагог организует работу с материалом",
    "дети выполняют практическое задание",
)


def _sections(text: str) -> dict[str, str]:
    matches = list(SECTION_RE.finditer(text))
    keys = [SECTION_NAMES[match.group(1).casefold()] for match in matches]
    if keys != ["goal", "tasks", "activities", "result"]:
        raise ValueError("нужны четыре части: цель, задачи, виды деятельности с ходом, результат")
    return {
        key: text[match.end():matches[index + 1].start() if index + 1 < len(matches) else len(text)].strip()
        for index, (key, match) in enumerate(zip(keys, matches))
    }


def _numbered_items(text: str) -> list[str]:
    matches = list(ITEM_RE.finditer(text))
    if not matches or len(matches) not in (2, 3) or text[:matches[0].start()].strip():
        raise ValueError("нужны 2–3 отдельные нумерованные пункта")
    return [
        text[match.end():matches[index + 1].start() if index + 1 < len(matches) else len(text)].strip()
        for index, match in enumerate(matches)
    ]


def _activity_fields(text: str) -> dict[str, str]:
    matches = list(ACTIVITY_RE.finditer(text))
    keys = [ACTIVITY_NAMES[match.group(1).casefold()] for match in matches]
    if keys != ["name", "materials", "procedure"]:
        raise ValueError("у каждой формы нужны название, материалы и ход")
    return {
        key: text[match.end():matches[index + 1].start() if index + 1 < len(matches) else len(text)].strip()
        for index, (key, match) in enumerate(zip(keys, matches))
    }


def check_cell(text: str) -> None:
    if any(phrase in text.casefold() for phrase in BOILERPLATE):
        raise ValueError("универсальная фраза заменяет конкретный ход работы")
    sections = _sections(text)
    if len(sections["goal"]) < 45 or len(sections["result"]) < 40:
        raise ValueError("цель или ожидаемый результат сформулированы слишком кратко")
    tasks = _numbered_items(sections["tasks"])
    if any(len(item) < 55 for item in tasks):
        raise ValueError("раскрой действие ребёнка и поддержку педагога в каждой задаче")
    activities = _numbered_items(sections["activities"])
    for index, activity in enumerate(activities, start=1):
        fields = _activity_fields(activity)
        if len(fields["name"]) < 12 or len(fields["materials"]) < 15:
            raise ValueError(f"форма {index}: назови конкретное действие и материалы")
        procedure = fields["procedure"]
        if len(procedure) < 150 or not ADULT_RE.search(procedure) or not CHILD_RE.search(procedure):
            raise ValueError(
                f"форма {index}: опиши последовательные действия педагога и детей с материалом"
            )


def check_table_shading(table) -> list[str]:
    """Reject colored fills, including shading inherited from table styles."""
    sources = [("таблица", table._tbl)]
    style = table.style
    seen = set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        sources.append((f"стиль таблицы {style.name}", style.element))
        style = style.base_style
    errors = []
    for label, element in sources:
        for shading in element.iter(qn("w:shd")):
            value = shading.get(qn("w:val"), "clear").lower()
            fill = shading.get(qn("w:fill"), "auto").upper()
            if value == "nil":
                continue
            themed = shading.get(qn("w:themeFill")) is not None
            if themed or fill not in {"AUTO", "FFFFFF"} or value != "clear":
                errors.append(f"{label}: запрещена цветная заливка; удалите её из ячеек и стиля таблицы")
                break
    return errors


def check_intro(document) -> list[str]:
    """Check the approved introduction before the annual-plan table."""
    elements = []
    for element in document.element.body:
        if element.tag == qn("w:tbl"):
            break
        elements.append(element)
    paragraphs = [p for p in document.paragraphs if p._p in elements and p.text.strip()]
    titles = {
        "Ұйымдастырылған іс-әрекеттің перспективалық жоспары":
            ["Мектепке дейінгі ұйым:", "тобы", "Балалардың жасы:", "оқу жылы", "Қамту кезеңі:"],
        "Перспективный план организованной деятельности":
            ["Дошкольная организация:", "Группа:", "Возраст детей:", "Учебный год:", "Период охвата:"],
        "Perspective plan of organized activities":
            ["Preschool organization:", "Group:", "Children's age:", "Academic year:", "Coverage period:"],
    }
    errors = []
    if not paragraphs or paragraphs[0].text.strip() not in titles:
        errors.append("верхний блок: требуется точный утверждённый заголовок плана")
    else:
        title = paragraphs[0]
        if title.alignment != WD_ALIGN_PARAGRAPH.CENTER:
            errors.append("верхний блок: заголовок должен быть по центру")
        lines = [p.text.strip() for p in paragraphs[1:]]
        position = 0
        for label in titles[title.text.strip()]:
            while position < len(lines) and label not in lines[position]:
                position += 1
            if position == len(lines):
                errors.append(f"верхний блок: отсутствует реквизит или нарушен порядок: {label}")
                break
            position += 1
    for element in elements:
        if any(br.get(qn("w:type")) == "page" for br in element.iter(qn("w:br"))):
            errors.append("верхний блок: нельзя отделять реквизиты от таблицы разрывом страницы")
            break
        if any(node.get(qn("w:val"), "1").lower() not in {"0", "false", "off"}
               for node in element.iter(qn("w:pageBreakBefore"))):
            errors.append("верхний блок: нельзя начинать реквизиты с новой страницы")
            break
        if any(True for _ in element.iter(qn("w:sectPr"))):
            errors.append("верхний блок: нельзя отделять реквизиты от таблицы разрывом раздела")
            break
    return errors


def audit(path: Path) -> list[str]:
    document = Document(path)
    errors: list[str] = []
    errors.extend(check_intro(document))
    for number, section in enumerate(document.sections, start=1):
        if section.orientation != WD_ORIENT.LANDSCAPE or section.page_width <= section.page_height:
            errors.append(f"раздел {number}: требуется альбомная ориентация страниц")
    if len(document.tables) != 1:
        errors.append("в перспективном плане должна быть одна таблица")
        return errors
    table = document.tables[0]
    errors.extend(check_table_shading(table))
    if len(table.columns) != 3:
        errors.append("таблица перспективного плана должна иметь три столбца")
        return errors
    for number, row in enumerate(table.rows[1:], start=2):
        subject = row.cells[1].text.strip() or "без названия деятельности"
        try:
            check_cell(row.cells[2].text.strip())
        except ValueError as exc:
            errors.append(f"строка {number}, {subject}: {exc}")
    return errors


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx", type=Path)
    args = parser.parse_args()
    errors = audit(args.docx)
    if errors:
        print(f"FAIL: {len(errors)} ошибок оформления или содержания", file=sys.stderr)
        for error in errors[:20]:
            print(f"- {error}", file=sys.stderr)
        if len(errors) > 20:
            print(f"... и ещё {len(errors) - 20}", file=sys.stderr)
        return 1
    print("PASS: альбомная ориентация, отсутствие цветной заливки и развёрнутые задачи с ходом работы проверены")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
