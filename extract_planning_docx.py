#!/usr/bin/env python3
"""Extract a verified planning DOCX into a lossless table payload for XLSX export."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table
from docx.oxml.ns import qn


COLUMNS = {"annual": 3, "weekly": 6, "daily": 2}


def extract(path: Path, mode: str) -> dict:
    document = Document(path)
    if len(document.tables) != 1:
        raise ValueError("В исходном DOCX должна быть ровно одна таблица плана.")
    table = document.tables[0]
    if len(table.columns) != COLUMNS[mode]:
        raise ValueError(f"Для {mode} требуется {COLUMNS[mode]} столбца(ов).")

    prefix, suffix = [], []
    past_table = False
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:tbl"):
            if past_table:
                raise ValueError("В исходном DOCX обнаружена лишняя таблица.")
            past_table = True
        elif child.tag == qn("w:p"):
            text = Paragraph(child, document._body).text.strip()
            if text:
                (suffix if past_table else prefix).append(text)

    rows = [[cell.text for cell in row.cells] for row in table.rows]
    if len(rows) < 2 or any(len(row) != COLUMNS[mode] for row in rows):
        raise ValueError("Таблица пуста или имеет неодинаковое число столбцов.")
    if any(len(value) > 32767 for row in rows for value in row):
        raise ValueError("Одна из ячеек DOCX превышает предел Excel в 32767 символов.")

    month_merges = []
    if mode == "annual":
        start = 1
        while start < len(rows):
            month = rows[start][0].strip()
            if not month:
                raise ValueError(f"У строки {start + 1} не указан месяц.")
            end = start + 1
            while end < len(rows) and rows[end][0].strip() == month:
                rows[end][0] = ""
                end += 1
            if end - start > 1:
                month_merges.append([start, end - 1])
            start = end

    return {"mode": mode, "prefix": prefix, "suffix": suffix,
            "rows": rows, "month_merges": month_merges}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--mode", choices=COLUMNS, required=True)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output-json", type=Path, required=True)
    args = parser.parse_args()
    payload = extract(args.input, args.mode)
    args.output_json.parent.mkdir(parents=True, exist_ok=True)
    args.output_json.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    print(f"Extracted {len(payload['rows']) - 1} data rows from {args.input.name}")


if __name__ == "__main__":
    main()
